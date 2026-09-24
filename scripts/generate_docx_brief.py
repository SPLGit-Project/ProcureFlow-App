import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_document():
    doc = docx.Document()

    # Configure 0.8 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Base style font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(30, 41, 59) # Slate 800

    # Color definitions
    NAVY = RGBColor(15, 41, 74)       # #0F294A
    BLUE = RGBColor(30, 64, 175)      # #1E40AF
    ACCENT_BLUE = RGBColor(37, 99, 235) # #2563EB
    MUTED = RGBColor(100, 116, 139)   # #64748B
    CRITICAL_RED = RGBColor(220, 38, 38)
    WARNING_AMBER = RGBColor(217, 119, 6)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_callout_border(cell, color_hex="2563EB"):
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>\n'
            f'  <w:left w:val="single" w:sz="36" w:space="0" w:color="{color_hex}"/>\n'
            f'  <w:top w:val="none"/>\n'
            f'  <w:right w:val="none"/>\n'
            f'  <w:bottom w:val="none"/>\n'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)

    # 1. Header with Logo & Title
    logo_path = 'docs/Logo Branding/APP-LOGOS/ProcureFlow-Title.png'
    if os.path.exists(logo_path):
        header_p = doc.add_paragraph()
        header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run_logo = header_p.add_run()
        run_logo.add_picture(logo_path, width=Inches(2.8))
        header_p.paragraph_format.space_after = Pt(4)

    title_p = doc.add_paragraph()
    title_run = title_p.add_run("Feature Release Brief: Dynamic Supplier Stock & 48-Hour Reservations")
    title_run.font.name = 'Calibri'
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = NAVY
    title_p.paragraph_format.space_after = Pt(2)
    title_p.paragraph_format.space_before = Pt(4)

    subtitle_p = doc.add_paragraph()
    sub_run = subtitle_p.add_run("Real-time available stock visibility, automated 48h fair-share reservations, and supplier stock analytics.")
    sub_run.font.name = 'Calibri'
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    sub_run.font.color.rgb = MUTED
    subtitle_p.paragraph_format.space_after = Pt(12)

    # 2. Metadata Information Table
    meta_table = doc.add_table(rows=2, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_headers = [
        ("Release Date", "24 September 2026"),
        ("Version", "Release v2.4 (Live)"),
        ("Audience", "Requisitioners, Approvers, Buyers"),
        ("Module", "Stock & Requisitions")
    ]

    for i, (label, val) in enumerate(meta_headers):
        col_idx = i
        c_top = meta_table.cell(0, col_idx)
        c_val = meta_table.cell(1, col_idx)

        set_cell_background(c_top, "0F294A")
        set_cell_background(c_val, "F8FAFC")
        set_cell_margins(c_top, top=80, bottom=60, left=120, right=120)
        set_cell_margins(c_val, top=80, bottom=80, left=120, right=120)

        # Border for cells
        for c in (c_top, c_val):
            tcPr = c._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>\n'
                f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>\n'
                f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>\n'
                f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>\n'
                f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>\n'
                f'</w:tcBorders>'
            )
            tcPr.append(tcBorders)

        p_lbl = c_top.paragraphs[0]
        p_lbl.paragraph_format.space_after = Pt(0)
        p_lbl.paragraph_format.space_before = Pt(0)
        r_lbl = p_lbl.add_run(label.upper())
        r_lbl.font.size = Pt(8)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = RGBColor(255, 255, 255)

        p_val = c_val.paragraphs[0]
        p_val.paragraph_format.space_after = Pt(0)
        p_val.paragraph_format.space_before = Pt(0)
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.5)
        r_val.font.bold = True
        r_val.font.color.rgb = NAVY

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 3. Executive Overview Callout
    callout_table = doc.add_table(rows=1, cols=1)
    callout_cell = callout_table.cell(0, 0)
    set_cell_background(callout_cell, "EFF6FF") # Soft Blue
    set_callout_border(callout_cell, "2563EB")
    set_cell_margins(callout_cell, top=140, bottom=140, left=180, right=180)

    p_call = callout_cell.paragraphs[0]
    p_call.paragraph_format.space_after = Pt(4)
    r_call_title = p_call.add_run("Executive Overview: Accurate Stock & Fair Inventory Allocation\n")
    r_call_title.font.bold = True
    r_call_title.font.size = Pt(11)
    r_call_title.font.color.rgb = BLUE

    r_call_body = p_call.add_run(
        "To prevent duplicate ordering and eliminate 'ghost stock' across SPL plants, ProcureFlow now tracks "
        "a live running total of available supplier stock. Approved requests place inventory into a temporary "
        "48-hour reservation holding pattern. If a Concur PO # is not linked within 48 hours, the order is "
        "automatically cancelled and the reserved stock is returned to the available pool for all other users."
    )
    r_call_body.font.size = Pt(10)
    r_call_body.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Helper for headings
    def add_section_heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = NAVY

        # Subtle bottom rule
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(
            f'<w:pBdr {nsdecls("w")}>\n'
            f'  <w:bottom w:val="single" w:sz="8" w:space="4" w:color="CBD5E1"/>\n'
            f'</w:pBdr>'
        )
        pPr.append(pBdr)

    # Helper for bullets
    def add_bullet(p_or_text, bold_prefix="", text=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.25)
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.bold = True
            r_bold.font.color.rgb = NAVY
        r_text = p.add_run(text)
        r_text.font.color.rgb = RGBColor(30, 41, 59)
        return p

    # --- Section 1: Core Functionality Breakdown ---
    add_section_heading("Key Enhancements & How They Work")

    add_bullet(
        None,
        "1. Real-Time Dynamic Running Stock: ",
        "Stock availability is no longer static between weekly supplier reports. ProcureFlow starts with the most "
        "recent supplier Stock-on-Hand (SOH) snapshot and continuously deducts both active reservations and committed "
        "orders in real time. The orderable quantity shown in the catalogue represents true, available inventory."
    )

    add_bullet(
        None,
        "2. Automated 48-Hour Stock Reservations: ",
        "The moment a requisition receives final approval, its items are placed into a 'Reserved' state. "
        "This temporarily earmarks the stock for that plant, preventing other sites from ordering the same inventory "
        "while procurement processes the purchase order."
    )

    add_bullet(
        None,
        "3. Live Real-Time Countdown Badges: ",
        "Every approved request features a visible countdown timer showing exactly how much time remains on the reservation. "
        "Color-coded urgency indicators keep buyers and requesters informed:"
    )

    # Countdown Urgency sub-table
    urgency_table = doc.add_table(rows=4, cols=3)
    urgency_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    urgency_headers = ["Urgency Tier", "Time Remaining", "Action Required"]
    for c_idx, h_text in enumerate(urgency_headers):
        cell = urgency_table.cell(0, c_idx)
        set_cell_background(cell, "0F294A")
        set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    urgency_data = [
        ("Normal (Blue)", "> 24 Hours", "Order is comfortably within the procurement processing window."),
        ("Warning (Amber)", "12 to 24 Hours", "Priority reminder: Concur PO # should be requested or entered promptly."),
        ("Critical (Red / Pulse)", "< 12 Hours", "Urgent action: Order will expire and auto-cancel if PO is not entered today.")
    ]

    for r_idx, (tier, time_left, action) in enumerate(urgency_data, start=1):
        for c_idx, text_val in enumerate([tier, time_left, action]):
            cell = urgency_table.cell(r_idx, c_idx)
            set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            # Add thin border
            tcPr = cell._tc.get_or_add_tcPr()
            tcBdr = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>\n'
                f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>\n'
                f'</w:tcBorders>'
            )
            tcPr.append(tcBdr)
            p = cell.paragraphs[0]
            r = p.add_run(text_val)
            r.font.size = Pt(9)
            if c_idx == 0:
                r.font.bold = True
                if "Red" in text_val:
                    r.font.color.rgb = CRITICAL_RED
                elif "Amber" in text_val:
                    r.font.color.rgb = WARNING_AMBER
                else:
                    r.font.color.rgb = BLUE

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_bullet(
        None,
        "4. Transition to Committed / In Delivery: ",
        "Once a buyer enters the Concur PO #, the order advances to 'Active (Awaiting Delivery)'. "
        "The inventory permanently shifts from 'Reserved' to 'Committed', remaining accounted for until delivery dockets "
        "and physical stock receipting are fully recorded."
    )

    add_bullet(
        None,
        "5. Automated Expiry & Cancellation: ",
        "If a Concur PO # is not linked before the 48-hour window closes, ProcureFlow's automated background engine "
        "cancels the request, returns the units to available stock, and issues an immediate automated notification to the requester."
    )

    # --- Section 2: Reporting & Analytics ---
    add_section_heading("New 'Stock & Reservations' Insights Report")

    p_rep = doc.add_paragraph()
    p_rep.paragraph_format.space_before = Pt(4)
    p_rep.paragraph_format.space_after = Pt(6)
    r_rep = p_rep.add_run(
        "A dedicated report has been added under "
    )
    r_rep_bold = p_rep.add_run("Reporting ➔ Supplier Insights ➔ Stock Reservations")
    r_rep_bold.font.bold = True
    r_rep_bold.font.color.rgb = BLUE
    p_rep.add_run(
        ", providing full transparency into inventory distribution across suppliers, active holds, and order pipeline health:"
    )

    add_bullet(None, "Executive KPI Metric Cards: ", "Real-time visibility into Total Net Orderable units ($ & units), Active 48h Reserved inventory, Committed in-delivery stock, Reservations Expiring Soon (<24h), and Auto-Cancelled Orders.")
    add_bullet(None, "Supplier Stock Breakdown Visual: ", "Stacked visual chart comparing Available vs Reserved vs Committed stock for every supplier, highlighting items under extreme reservation pressure (>70% held).")
    add_bullet(None, "Live 48h Reservation Queue: ", "Interactive queue allowing procurement officers to filter orders by urgency tier, inspect time remaining down to the minute, and jump directly into the PO with one click.")
    add_bullet(None, "Auto-Cancellation Audit Log: ", "Complete audit log recording every expired requisition, units returned to inventory, and expiration timestamps.")
    add_bullet(None, "Comprehensive CSV Export: ", "Exports the full dataset including internal SKUs, supplier codes, baseline SOH, unit prices, effective stock, and reservation pressure ratings.")

    # --- Section 3: Role-by-Role Action Matrix ---
    add_section_heading("What This Means for You: Role-by-Role Guidance")

    role_table = doc.add_table(rows=4, cols=3)
    role_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    role_headers = ["Role", "Key Benefits", "Action Items & Best Practices"]
    for c_idx, h_text in enumerate(role_headers):
        cell = role_table.cell(0, c_idx)
        set_cell_background(cell, "0F294A")
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    role_data = [
        (
            "Requisitioners\n(Site / Plant)",
            "• Immediate confirmation that approved items are held exclusively for your plant.\n• Accurate catalogue stock prevents wasted requests.",
            "• Check the available stock pill when browsing the catalogue.\n• Once approved, track your order's 48h countdown badge.\n• Promptly check your email / notifications if your order is nearing 12h remaining."
        ),
        (
            "Procurement Officers\n(Buyers)",
            "• Prevents stock being double-allocated to competing plants.\n• Prioritized queue flags orders needing immediate Concur PO # generation.",
            "• Monitor the 'Pending Concur PO #' filter in your Task Drawer.\n• Prioritize orders showing amber or red countdown badges.\n• Enter the Concur PO # promptly to secure inventory into 'Active' delivery status."
        ),
        (
            "Approvers &\nPlant Managers",
            "• Confidence that approved requests correspond to confirmed available supplier stock.\n• Automated cleanup of stale or abandoned orders.",
            "• Approve requisitions promptly so buyers receive the full 48-hour procurement window.\n• Review the Stock Reservations report to monitor supplier stock trends and site demand."
        )
    ]

    for r_idx, (role, benefits, actions) in enumerate(role_data, start=1):
        for c_idx, text_val in enumerate([role, benefits, actions]):
            cell = role_table.cell(r_idx, c_idx)
            set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            tcPr = cell._tc.get_or_add_tcPr()
            tcBdr = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>\n'
                f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>\n'
                f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>\n'
                f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>\n'
                f'</w:tcBorders>'
            )
            tcPr.append(tcBdr)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text_val)
            r.font.size = Pt(9)
            if c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = NAVY

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- Section 4: Frequently Asked Questions ---
    add_section_heading("Frequently Asked Questions (FAQ)")

    faqs = [
        (
            "Q: What happens if an order is auto-cancelled after 48 hours?",
            "A: The request moves to 'Cancelled' status and its reserved items are immediately returned to available supplier stock for any plant to request. The requester receives an automated system alert with full order details. If the items are still required, a new requisition can be raised."
        ),
        (
            "Q: Can a reservation be extended beyond 48 hours?",
            "A: The 48-hour window is intentionally strict to enforce fair stock allocation across all national sites and prevent stock locking. Procurement teams should prioritize orders displaying Warning (<24h) or Critical (<12h) badges."
        ),
        (
            "Q: How does this interact with weekly supplier inventory spreadsheets?",
            "A: ProcureFlow seamlessly reconciles weekly supplier inventory uploads against active reservations. New weekly reports reset the baseline SOH, while unfulfilled committed orders and active 48h reservations continue to deduct dynamically."
        ),
        (
            "Q: What happens when an order is partially delivered?",
            "A: Only unreceived quantities remain committed. For example, if 50 units were ordered and 30 have been receipted, only the remaining 20 units continue to be deducted from available supplier stock."
        )
    ]

    for q, a in faqs:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.keep_with_next = True
        r_q = p_q.add_run(q)
        r_q.font.bold = True
        r_q.font.size = Pt(10)
        r_q.font.color.rgb = BLUE

        p_a = doc.add_paragraph()
        p_a.paragraph_format.space_before = Pt(0)
        p_a.paragraph_format.space_after = Pt(6)
        r_a = p_a.add_run(a)
        r_a.font.size = Pt(9.5)
        r_a.font.color.rgb = RGBColor(51, 65, 85)

    # --- Section 5: Support & Contact ---
    add_section_heading("Questions & Support")
    p_sup = doc.add_paragraph()
    p_sup.paragraph_format.space_before = Pt(4)
    p_sup.paragraph_format.space_after = Pt(6)
    p_sup.add_run(
        "For system support, supplier stock discrepancies, or questions regarding dynamic reservations:\n"
    )
    add_bullet(None, "ProcureFlow Support Desk: ", "Submit a ticket via the in-app Help Guide & Feedback drawer.")
    add_bullet(None, "Procurement Systems Team: ", "Contact your central procurement category manager or buyer.")

    # Save document
    output_path = 'docs/ProcureFlow_Update_Brief_Dynamic_Supplier_Stock_and_Reservations.docx'
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == '__main__':
    create_document()
